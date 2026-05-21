# -*- coding: utf-8 -*-
"""Genera Manual-Despliegue-RRHH.docx — estilo documento SENA (sangría, negritas)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent.parent / "docs" / "Manual-Despliegue-RRHH.docx"

INTEGRANTES = [
    "Ronaldo Stiven Franco Durán",
    "Briyid Tatiana Cruz Molina",
    "Willington Guzmán Arias",
    "Angela Tatiana Gonzalez Pinto",
]
INSTRUCTOR = "Ing. Astrid Segura"
FICHA = "2826074"


def _font(run, size=11, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold


def parrafo(doc, texto, sangria=True, centrado=False, negrita=False, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrado else WD_ALIGN_PARAGRAPH.JUSTIFY
    if sangria and not centrado:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(texto)
    _font(r, size=size, bold=negrita)
    return p


def titulo_seccion(doc, texto, nivel=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(texto)
    _font(r, size=14 if nivel == 1 else 12, bold=True)
    return p


def nota(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run("Nota: ")
    _font(r1, bold=True)
    r2 = p.add_run(texto)
    _font(r2)


def paso(doc, numero, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"{numero}. ")
    _font(r1, bold=True)
    r2 = p.add_run(texto)
    _font(r2)


def bloque_codigo(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(texto)
    r.font.name = "Consolas"
    r.font.size = Pt(9)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(3)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(2.5)

    # —— Portada ——
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Manual de despliegue")
    _font(r, size=18, bold=True)

    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(
        "Sistema de información para la gestión de Recursos Humanos\n"
        "(PROYECTO RRHH — Talent Sphere)"
    )
    _font(r2, size=13, bold=True)

    doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run("Aprendices:\n")
    _font(r3, size=12, bold=True)
    for nombre in INTEGRANTES:
        rn = p3.add_run(f"{nombre}\n")
        _font(rn, size=12)

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p4.add_run(f"Instructor:\n{INSTRUCTOR}\n")
    _font(r4, size=12, bold=False)
    r4b = p4.add_run(
        "Tecnólogo en Análisis y Desarrollo de Software\n"
        f"Ficha {FICHA}\n"
        "Centro de Diseño y Metrología\n"
        "Servicio Nacional de Aprendizaje — SENA\n"
        "Bogotá D.C. — 2026"
    )
    _font(r4b, size=12)

    doc.add_page_break()

    # —— Objetivo ——
    titulo_seccion(doc, "Objetivo")
    parrafo(
        doc,
        "El propósito de este manual es describir, de forma ordenada, los pasos para instalar, "
        "configurar y dejar en operación el sistema de Recursos Humanos del proyecto. La solución "
        "combina un frontend en React publicado en GitHub Pages y una API en Laravel desplegada en "
        "Azure Container Apps, con base de datos MySQL en Aiven Cloud.",
    )
    parrafo(
        doc,
        "La guía está pensada para que cualquier integrante del equipo pueda repetir el despliegue "
        "en otro equipo o ambiente, siguiendo el mismo orden: base de datos, backend, pruebas de API "
        "y por último la publicación del frontend.",
    )

    # —— 1. Arquitectura ——
    titulo_seccion(doc, "1. Arquitectura general del sistema")
    parrafo(
        doc,
        "La tabla siguiente resume dónde corre cada componente en producción:",
    )

    table = doc.add_table(rows=6, cols=3)
    table.style = "Table Grid"
    headers = ["Capa", "Tecnología", "Ubicación"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        _font(run, bold=True)

    datos = [
        ("Interfaz de usuario", "React 19, Vite 6, HashRouter", "GitHub Pages"),
        ("API REST", "Laravel (PHP 8.2) en contenedor Docker", "Azure Container Apps"),
        ("Registro de imágenes", "Dockerfile del repo Laravel", "Azure Container Registry"),
        ("Datos", "MySQL con conexión SSL", "Aiven Cloud"),
        ("Correo", "SMTP", "Variables en Container App"),
    ]
    for ri, fila in enumerate(datos, 1):
        for ci, val in enumerate(fila):
            table.rows[ri].cells[ci].text = val

    parrafo(doc, "")
    parrafo(doc, "Direcciones de referencia en producción:", negrita=True, sangria=False)
    bloque_codigo(
        doc,
        "API:  https://mi-api.lemonforest-f9c17ad2.eastus.azurecontainerapps.io/api/v1\n"
        "Login: https://mi-api.lemonforest-f9c17ad2.eastus.azurecontainerapps.io/api/v1/login\n"
        "Front: https://ronald-f18.github.io/PROYECTO-REACT-RRHH/#/login",
    )
    parrafo(
        doc,
        "El flujo de datos es el siguiente: el usuario abre el sitio en GitHub Pages; el navegador "
        "consume la API en Azure con token Bearer; Laravel consulta MySQL en Aiven.",
    )

    # —— 2. Frontend ——
    titulo_seccion(doc, "2. Despliegue del frontend (GitHub Pages)")
    titulo_seccion(doc, "2.1 Repositorio y compilación", 2)
    parrafo(
        doc,
        "El código del cliente está en el repositorio PROYECTO-REACT-RRHH. Para generar el sitio "
        "estático se ejecuta npm run build, que deja los archivos en la carpeta dist/. La publicación "
        "se hace con npm run deploy, el cual sube dist/ a la rama gh-pages mediante gh-pages.",
    )
    titulo_seccion(doc, "2.2 Rutas y configuración de Vite", 2)
    parrafo(
        doc,
        "GitHub Pages sirve el proyecto en una subcarpeta. Por eso en vite.config.js debe quedar "
        "base: '/PROYECTO-REACT-RRHH/'. El enrutador es HashRouter, de modo que las rutas se ven "
        "como #/login, #/dashboard o #/empleados.",
    )
    titulo_seccion(doc, "2.3 Variables de entorno del frontend", 2)
    parrafo(
        doc,
        "Antes del build de producción se configura el archivo .env.production. La variable "
        "principal es VITE_API_URL, que debe apuntar a la API en Azure sin barra al final.",
    )
    bloque_codigo(
        doc,
        "VITE_API_URL=https://mi-api.lemonforest-f9c17ad2.eastus.azurecontainerapps.io/api/v1",
    )
    nota(
        doc,
        "Las variables VITE_* quedan grabadas en el JavaScript al compilar. Si cambia la URL del API, "
        "hay que volver a ejecutar npm run build y npm run deploy.",
    )
    titulo_seccion(doc, "2.4 Comunicación con la API (autenticación)", 2)
    parrafo(
        doc,
        "En producción el frontend no envía cookies entre dominios (withCredentials: false). "
        "El inicio de sesión se hace con POST /api/v1/login enviando email_usuario y "
        "contrasena_usuario. La respuesta trae access_token, token_type, user y role. Ese token "
        "se guarda en localStorage y se envía en las peticiones siguientes como "
        "Authorization: Bearer {token}.",
    )
    parrafo(
        doc,
        "En Laravel debe estar permitido el origen https://ronald-f18.github.io en la configuración "
        "de CORS, con los encabezados Authorization, Content-Type y Accept.",
    )
    titulo_seccion(doc, "2.5 Endpoints que usa la interfaz", 2)
    parrafo(
        doc,
        "El panel principal consume un solo endpoint de resumen: GET /api/v1/dashboard/resumen. "
        "Los módulos de empleados y contratos usan paginación con page y per_page; la respuesta "
        "incluye data y meta. El resto de módulos sigue usando listados completos hasta que se "
        "amplíe la paginación en backend.",
    )
    titulo_seccion(doc, "2.6 Pasos para publicar el frontend", 2)
    paso(doc, 1, "Clonar el repositorio React e instalar dependencias (npm install).")
    paso(doc, 2, "Crear o editar .env.production con VITE_API_URL correcto.")
    paso(doc, 3, "Ejecutar npm run build.")
    paso(doc, 4, "Ejecutar npm run deploy con una cuenta GitHub con permisos sobre el repo.")
    paso(
        doc,
        5,
        "En GitHub → Settings → Pages, confirmar origen: rama gh-pages, carpeta raíz (/).",
    )
    paso(
        doc,
        6,
        "Abrir la URL del sitio y revisar en las herramientas de desarrollador (F12) que las "
        "peticiones salgan hacia azurecontainerapps.io.",
    )

    # —— 3. Aiven ——
    titulo_seccion(doc, "3. Base de datos en Aiven")
    parrafo(
        doc,
        "MySQL no va dentro del contenedor de la API. Se contrata como servicio en Aiven. Desde "
        "la consola se crea el servicio, se descarga el certificado CA y se guarda en la raíz del "
        "proyecto Laravel con el nombre ca.pem.",
    )
    paso(doc, 1, "Registrar el proyecto en https://console.aiven.io y crear servicio MySQL.")
    paso(doc, 2, "Anotar host, puerto, base de datos, usuario y contraseña para el archivo .env.")
    paso(doc, 3, "Colocar ca.pem en la raíz del backend (se copia en la imagen Docker).")
    paso(doc, 4, "Definir MYSQL_ATTR_SSL_CA=/var/www/html/ca.pem en las variables del contenedor.")
    nota(
        doc,
        "La base está en la nube y la API en Azure East US; la latencia de red es normal. "
        "Conviene revisar la lista de IP permitidas en Aiven si el servicio lo exige.",
    )

    # —— 4. Azure ——
    titulo_seccion(doc, "4. Despliegue del backend en Azure")
    titulo_seccion(doc, "4.1 Recursos utilizados", 2)
    parrafo(doc, "En la suscripción Azure for Students, grupo de recursos api (East US), se usan:")
    paso(doc, 1, "Container Registry: acrronaldsena2026.")
    paso(doc, 2, "Container App: mi-api (entorno mi-api-env).")
    paso(doc, 3, "Imagen de referencia actual: acrronaldsena2026.azurecr.io/mi-api:v7.")
    paso(doc, 4, "Log Analytics: workspace-apiHCSK.")
    titulo_seccion(doc, "4.2 Requisitos en el equipo de despliegue", 2)
    paso(doc, 1, "Docker Desktop encendido.")
    paso(doc, 2, "Azure CLI con sesión iniciada (az login) y suscripción Azure for Students.")
    paso(doc, 3, "Archivo .env local con credenciales (no subir a Git).")
    paso(doc, 4, "Archivo ca.pem en la raíz del proyecto Laravel.")
    titulo_seccion(doc, "4.3 Actualización habitual (scripts del repo)", 2)
    parrafo(
        doc,
        "Desde PowerShell, en la carpeta PROYECTO-RRHH, con política de ejecución permitida "
        "en la sesión:",
    )
    bloque_codigo(
        doc,
        "Set-ExecutionPolicy -Scope Process -ExecutionPolicy Bypass -Force\n"
        "cd C:\\Users\\APRENDIZ\\Documents\\PROYECTO-RRHH\n"
        "az account set --subscription \"Azure for Students\"\n"
        ".\\scripts\\azure-push-docker.ps1 -ImageTag v7\n"
        ".\\scripts\\azure-deploy.ps1 -SkipBuild -ImageTag v7\n"
        ".\\scripts\\azure-status.ps1\n"
        ".\\scripts\\azure-test-login.ps1",
    )
    parrafo(
        doc,
        "El primer comando construye y sube la imagen al registro. El segundo actualiza la "
        "Container App para usar ese tag. Al arrancar el contenedor se ejecutan las migraciones "
        "con php artisan migrate --force.",
    )
    nota(doc, "En cada entrega nueva se cambia el tag (v8, v9, etc.) para no sobrescribir versiones anteriores.")
    titulo_seccion(doc, "4.4 Creación inicial (solo la primera vez)", 2)
    bloque_codigo(doc, "az group create --name api --location eastus")
    bloque_codigo(doc, "az acr create --resource-group api --name acrronaldsena2026 --sku Basic")
    bloque_codigo(doc, "az acr update --name acrronaldsena2026 --admin-enabled true")
    titulo_seccion(doc, "4.5 Variables de entorno del backend", 2)
    parrafo(
        doc,
        "Los valores sensibles no deben publicarse en documentos ni en repositorios. "
        "En el manual solo se listan los nombres de variable:",
    )
    vars = [
        "APP_KEY, APP_URL, APP_DEBUG=false",
        "DB_HOST, DB_PORT, DB_DATABASE, DB_USERNAME, DB_PASSWORD",
        "MYSQL_ATTR_SSL_CA=/var/www/html/ca.pem",
        "JWT_SECRET",
        "MAIL_MAILER, MAIL_HOST, MAIL_PORT, MAIL_USERNAME, MAIL_PASSWORD",
        "FRONTEND_URL=https://ronald-f18.github.io",
        "SANCTUM_STATEFUL_DOMAINS (localhost y ronald-f18.github.io)",
    ]
    for v in vars:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.25)
        r = p.add_run("• " + v)
        _font(r)
    titulo_seccion(doc, "4.6 Ajustes recomendados en Container App", 2)
    paso(doc, 1, "Réplicas mínimas en 1 para reducir la espera de la primera petición.")
    paso(doc, 2, "Puerto de entrada del contenedor: 8000.")
    paso(doc, 3, "Revisar logs con .\\scripts\\azure-logs.ps1 si hay errores al iniciar.")

    # —— 5. Orden ——
    titulo_seccion(doc, "5. Orden sugerido de despliegue completo")
    paso(doc, 1, "Tener la base Aiven creada y el archivo ca.pem en el proyecto.")
    paso(doc, 2, "Subir imagen Docker a ACR (azure-push-docker.ps1).")
    paso(doc, 3, "Desplegar revisión en Container App (azure-deploy.ps1).")
    paso(doc, 4, "Probar login y endpoints con azure-test-login.ps1 o Postman.")
    paso(doc, 5, "Compilar y publicar frontend con VITE_API_URL apuntando a la API.")
    paso(doc, 6, "Probar login y dashboard desde GitHub Pages.")

    # —— 6. Checklist ——
    titulo_seccion(doc, "6. Lista de verificación después del despliegue")
    checks = [
        "Container App en estado Running.",
        "POST /login devuelve access_token.",
        "GET /dashboard/resumen responde 200 con Bearer.",
        "GET /empleados?page=1 devuelve data y meta.",
        "Sin errores CORS desde GitHub Pages.",
        "Frontend publicado con la URL de API correcta.",
    ]
    for c in checks:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        r = p.add_run("☐ " + c)
        _font(r)

    # —— 7. Rollback ——
    titulo_seccion(doc, "7. Volver a una versión anterior")
    parrafo(
        doc,
        "Si una revisión nueva presenta fallos, se puede desplegar de nuevo un tag anterior "
        "(por ejemplo v6) con azure-deploy.ps1 -SkipBuild -ImageTag v6, o desde el portal de "
        "Azure mover el tráfico a la revisión previa en Container App → Revisiones.",
    )

    # —— 8. Problemas frecuentes ——
    titulo_seccion(doc, "8. Problemas frecuentes y qué revisar")
    problemas = [
        (
            "Error 401 en rutas protegidas",
            "Comprobar que el front envía Bearer y que el token no expiró; repetir login.",
        ),
        (
            "CORS bloqueado en el navegador",
            "Verificar config/cors.php y FRONTEND_URL con https://ronald-f18.github.io.",
        ),
        (
            "Primera carga muy lenta",
            "Cold start o BD lejana; subir minReplicas a 1 y usar dashboard/resumen en el panel.",
        ),
        (
            "Error de conexión a MySQL",
            "Revisar ca.pem, credenciales en secretos de Container App y firewall de Aiven.",
        ),
        (
            "Script PowerShell no ejecuta",
            "Set-ExecutionPolicy -Scope Process Bypass o ejecutar con powershell -ExecutionPolicy Bypass -File.",
        ),
    ]
    for sintoma, accion in problemas:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.space_after = Pt(6)
        r1 = p.add_run(sintoma + ". ")
        _font(r1, bold=True)
        r2 = p.add_run(accion)
        _font(r2)

    # —— 9. Seguridad ——
    titulo_seccion(doc, "9. Seguridad")
    parrafo(
        doc,
        "No incluir en Git ni en entregables públicos el archivo .env, contraseñas de base de datos, "
        "APP_KEY, JWT_SECRET ni contraseñas de correo. Si alguna credencial se compartió por error, "
        "conviene rotarla en Aiven, Azure y Gmail. Los secretos deben gestionarse desde Azure Container "
        "Apps o, en un entorno más formal, Azure Key Vault.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Generado: {OUT}")


if __name__ == "__main__":
    build()
